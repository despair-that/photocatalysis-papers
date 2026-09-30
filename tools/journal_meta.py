"""Build data/journal_meta.json from up to three sources.

Sources (any subset; maps not mentioned are kept from the existing file):
    --sjr <scimagojr.csv|.zip>   SJR indicator, SJR quartile, H-index
                                 (scimagojr.com, behind Cloudflare — download
                                 in a real browser once a year)
    --cas <中科院分区表.xlsx>     CAS quartile ("2025分区") + Top flag; sheet
                                 "2023 vs 2025分区对比", columns
                                 期刊名称/2025分区/2023分区/Top/Open Access
    --if  <IF....xlsx>           Journal impact factors, columns
                                 Journal / IF_2022 .. IF_2026; the latest
                                 non-empty column wins and its year is stored

Output: data/journal_meta.json with by_issn/by_title (SJR), by_cas, by_if.
Keys are normalized titles (lowercase, punctuation stripped); matching at
runtime is ISSN first, then title, then unique first-3-word prefix.
"""

import argparse
import csv
import io
import json
import re
import sys
import zipfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "journal_meta.json"
_NON_ALNUM = re.compile(r"[^a-z0-9 ]+")
IF_COLUMNS = ("IF_2026", "IF_2025", "IF_2024", "IF_2023", "IF_2022")


def normalize_title(name: str) -> str:
    words = _NON_ALNUM.sub(" ", str(name or "").lower()).split()
    if words and words[0] == "the":
        words = words[1:]
    return " ".join(words)


def fnum(value):
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def read_sjr(path: str):
    src = Path(path)
    if src.suffix.lower() == ".zip":
        with zipfile.ZipFile(src) as archive:
            name = next(n for n in archive.namelist()
                        if n.lower().endswith(".csv"))
            raw = archive.read(name)
    else:
        raw = src.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        sys.exit("could not decode the SJR export")
    by_issn: dict = {}
    by_title: dict = {}
    for row in csv.DictReader(io.StringIO(text), delimiter=";"):
        title = (row.get("Title") or "").strip()
        if not title:
            continue
        quartile = (row.get("SJR Best Quartile") or "").strip().upper()
        if quartile not in ("Q1", "Q2", "Q3", "Q4"):
            quartile = ""
        h_index = fnum(row.get("H index"))
        entry = {"q": quartile, "s": fnum(row.get("SJR")),
                 "h": int(h_index) if h_index else None}
        for code in re.split(r"[,; ]+", (row.get("Issn") or "")):
            code = code.strip().replace("-", "").lower()
            if len(code) >= 8:
                by_issn[code] = entry
        key = normalize_title(title)
        if key:
            by_title[key] = entry
    return by_issn, by_title


def read_cas(path: str):
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True)
    sheet = workbook.worksheets[0]
    by_cas: dict = {}
    conflicts = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        title = str(row[0] or "").strip()
        if not title:
            continue
        quartile = str(row[1] or "").strip()
        if quartile not in ("1", "2", "3", "4"):
            continue
        top = str(row[3] or "").strip() == "是"
        key = normalize_title(title)
        entry = {"cas": quartile, "top": top}
        if key in by_cas and by_cas[key] != entry:
            conflicts += 1
            continue
        by_cas[key] = entry
    workbook.close()
    return by_cas, conflicts


def read_if(path: str):
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True)
    sheet = workbook.worksheets[0]
    header = [c for c in next(sheet.iter_rows(max_row=1, values_only=True))]
    by_if: dict = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        title = str(row[0] or "").strip()
        if not title:
            continue
        values = dict(zip(header, row))
        for column in IF_COLUMNS:
            factor = fnum(values.get(column))
            if factor is not None:
                by_if[normalize_title(title)] = {
                    "if": factor, "year": int(column.split("_")[1])}
                break
    workbook.close()
    return by_if


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sjr")
    parser.add_argument("--cas")
    parser.add_argument("--if", dest="if_file")
    args = parser.parse_args()
    if not (args.sjr or args.cas or args.if_file):
        sys.exit("nothing to do: pass --sjr / --cas / --if")

    meta = {}
    if OUT.exists():
        meta = json.loads(OUT.read_text(encoding="utf-8"))

    if args.sjr:
        meta["by_issn"], meta["by_title"] = read_sjr(args.sjr)
        print(f"SJR: {len(meta['by_title'])} titles, "
              f"{len(meta['by_issn'])} ISSNs")
    if args.cas:
        meta["by_cas"], conflicts = read_cas(args.cas)
        print(f"CAS: {len(meta['by_cas'])} journals "
              f"({conflicts} conflicting duplicates skipped)")
    if args.if_file:
        meta["by_if"] = read_if(args.if_file)
        print(f"IF: {len(meta['by_if'])} journals")

    meta["updated"] = date.today().isoformat()
    meta["source"] = "SCImago (scimagojr.com) + 中科院分区表2025 + Journal IF"
    OUT.write_text(json.dumps(meta, ensure_ascii=False, separators=(",", ":")),
                   encoding="utf-8")
    print(f"journal_meta.json written: "
          f"{OUT.stat().st_size / 1048576:.1f} MB")


if __name__ == "__main__":
    main()
